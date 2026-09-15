from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

from openpyxl import load_workbook

from .csv_repository import ComponentRepository
from .models import Component, MCS
from .parsing import first_number, normalize_application, parse_power_w, parse_range


SHEET_CATEGORIES = {
    "Beamformer_BFIC": "BFIC",
    "FEM_PA_LNA_Switch": "FEM",
    "Converter_Mixer": "MIXER",
    "PLL_Synthesizer": "PLL",
    "DataConverter_ADC_DAC": "CONVERTER",
}

FIELD_MAP = {
    "Number of Channels (elements)": "number_channels",
    "Number of Channels": "number_channels",
    "Number of Beams": "number_beams",
    "Gain (single channel)": "gain_db",
    "Channel Bandwidth": "channel_bandwidth_ghz",
    "Output P1dB (OP1dB)": "output_p1db_dbm",
    "Output P1dB / P3dB": "output_p1db_dbm",
    "Output P1dB": "output_p1db_dbm",
    "Psat per Channel": "psat_dbm",
    "Saturated Power (Psat)": "psat_dbm",
    "OIP3": "oip3_dbm",
    "Noise Figure (NF)": "noise_figure_db",
    "Input P1dB (IP1dB)": "input_p1db_dbm",
    "RMS Phase Error": "rms_phase_error_deg",
    "RMS Gain Error": "rms_gain_error_db",
    "EVM (per modulation)": "evm_percent",
    "Linear Gain": "pa_gain_db",
    "Efficiency (PAE)": "pa_pae_percent",
    "Gain": "lna_gain_db",
    "Input IP3 (IIP3)": "lna_iip3_dbm",
    "Insertion Loss": "insertion_loss_db",
    "Isolation": "isolation_db",
    "Power Handling": "power_handling_dbm",
    "Conversion Gain / Loss": "conversion_gain_db",
    "LO drive level": "lo_drive_dbm",
    "Output Power": "output_power_dbm",
    "Phase Noise @ 1 kHz": "phase_noise_1khz",
    "Phase Noise @ 10 kHz": "phase_noise_10khz",
    "Phase Noise @ 100 kHz": "phase_noise_100khz",
    "Phase Noise @ 1 MHz": "phase_noise_1mhz",
    "Resolution": "resolution_bit",
    "Sample Rate": "sample_rate_gsps",
    "Analog Input Bandwidth": "analog_bandwidth_ghz",
    "SFDR": "sfdr_dbc",
    "SNR / SINAD": "snr_db",
    "ENOB": "enob_bit",
}

RANGE_MAP = {
    "Operating Frequency Range": ("freq_min_ghz", "freq_max_ghz"),
    "RF Frequency Range": ("rf_min_ghz", "rf_max_ghz"),
    "IF Frequency Range": ("if_min_ghz", "if_max_ghz"),
    "LO Frequency Range": ("lo_min_ghz", "lo_max_ghz"),
    "Output Frequency Range": ("output_freq_min_ghz", "output_freq_max_ghz"),
    "Operating Temperature": ("operating_temp_min_c", "operating_temp_max_c"),
}

REFERENCE_PRODUCT_URLS = {
    "Stampede2731 (HP)": "https://www.sivers-semiconductors.com/wireless/stampede2731-2/",
    "Stampede2731 (MP/LP)": "https://www.sivers-semiconductors.com/wireless/stampede2731lp-2/",
}

REFERENCE_PRODUCT_IMAGES = {
    "Stampede2731 (HP)": "https://www.sivers-semiconductors.com/wp-content/uploads/2025/11/Stampede2731-1.png",
}


def _rows_by_parameter(ws) -> Dict[str, int]:
    return {str(ws.cell(row, 2).value).strip(): row for row in range(1, ws.max_row + 1) if ws.cell(row, 2).value}


def _actual_model_columns(ws) -> Iterable[int]:
    for column in range(5, ws.max_column + 1):
        header = ws.cell(4, column).value
        if header and str(header).strip() and str(header).strip() != "예시 (Example)":
            yield column


def import_components(path: Path, repository: ComponentRepository) -> Tuple[int, List[str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    imported = 0
    warnings: List[str] = []
    for sheet_name, category in SHEET_CATEGORIES.items():
        if sheet_name not in workbook.sheetnames:
            warnings.append(f"시트 없음: {sheet_name}")
            continue
        ws = workbook[sheet_name]
        rows = _rows_by_parameter(ws)
        for column in _actual_model_columns(ws):
            def raw(parameter: str) -> Any:
                row = rows.get(parameter)
                return ws.cell(row, column).value if row else None

            part_no = raw("Part No.")
            manufacturer = raw("Manufacturer")
            if not part_no or not manufacturer:
                continue
            application = normalize_application(raw("Application (Terminal / Payload)"))
            function_label = next((key for key in rows if key.startswith("Function (")), "")
            function = str(raw(function_label) or "")
            freq_label = "Operating Frequency Range" if "Operating Frequency Range" in rows else "RF Frequency Range"
            freq_min, freq_max = parse_range(raw(freq_label))
            supply = first_number(raw("Supply Voltage"))
            power_raw = raw("Power Consumption") or raw("Total DC Power")
            power_unit = str(ws.cell(rows.get("Power Consumption", rows.get("Total DC Power", 1)), 3).value or "")
            specs: Dict[str, Any] = {}
            common_values: Dict[str, Any] = {}
            for label, (low_key, high_key) in RANGE_MAP.items():
                low, high = parse_range(raw(label))
                if label == "Operating Frequency Range":
                    continue
                if label == "Operating Temperature":
                    common_values[low_key], common_values[high_key] = low, high
                else:
                    specs[low_key], specs[high_key] = low, high
            for label, key in FIELD_MAP.items():
                value = raw(label)
                if value is None or str(value).strip() in {"", "—", "TBD"}:
                    continue
                if key == "power_handling_dbm" and value is not None:
                    number = first_number(value)
                    unit = str(ws.cell(rows[label], 3).value or "").lower()
                    if number is not None and unit == "w":
                        import math
                        number = 30.0+10.0 * math.log10(number)
                    specs[key] = number
                elif key == "lo_drive_dbm":
                    low, high = parse_range(value)
                    specs["lo_drive_min_dbm"], specs["lo_drive_max_dbm"] = low, high
                else:
                    specs[key] = first_number(value)
            component = Component(
                component_id=None,
                category=category,
                manufacturer=str(manufacturer).strip(),
                part_no=str(part_no).strip(),
                application=application,
                function=function.strip(),
                process=str(raw(next((x for x in rows if x.startswith("Process (")), "")) or "") or None,
                freq_min_ghz=freq_min,
                freq_max_ghz=freq_max,
                supply_voltage_v=supply,
                power_consumption_w=parse_power_w(power_raw, power_unit),
                package=str(raw("Package") or raw("Package / Size") or "") or None,
                operating_temp_min_c=common_values.get("operating_temp_min_c"),
                operating_temp_max_c=common_values.get("operating_temp_max_c"),
                product_url=REFERENCE_PRODUCT_URLS.get(str(part_no).strip()),
                image_url=REFERENCE_PRODUCT_IMAGES.get(str(part_no).strip()),
                source_file=Path(path).name,
                note=f"Excel slot: {ws.cell(4, column).value}",
                specs={key: value for key, value in specs.items() if value is not None},
            )
            repository.upsert(component)
            imported += 1
    return imported, warnings


def import_mcs(path: Path, repository: ComponentRepository) -> int:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if "MCS" not in workbook.sheetnames:
        raise ValueError("MCS 시트를 찾을 수 없습니다.")
    ws = workbook["MCS"]
    rows: List[MCS] = []
    for row in range(4, ws.max_row + 1):
        index = ws.cell(row, 2).value
        order = ws.cell(row, 3).value
        minimum = ws.cell(row, 4).value
        maximum = ws.cell(row, 5).value
        efficiency = ws.cell(row, 7).value
        rate = ws.cell(row, 8).value
        if all(isinstance(x, (int, float)) for x in (index, order, minimum, maximum, efficiency)):
            rows.append(MCS(int(index), int(order), float(minimum), float(maximum), float(efficiency), float(rate) / 1024.0 if rate else None))
    repository.replace_mcs(rows)
    return len(rows)


COMMON_CSV_FIELDS = {
    "component_id","category","manufacturer","part_no","application","function","process",
    "freq_min_ghz","freq_max_ghz","supply_voltage_v","power_consumption_w",
    "package","operating_temp_min_c","operating_temp_max_c","product_url",
    "datasheet_url","datasheet_revision","datasheet_page","data_origin","retrieved_at",
    "review_status","slot","note",
    "source_file","source_date","extraction_method","extraction_evidence","image_url",
}


def _csv_number(value: Any):
    return first_number(value) if value not in (None,"") else None


def component_from_csv_row(row: dict, source_file: str = "csv") -> Component:
    category = row.get("category", "").upper()
    function = row.get("function", "")
    if category == "FEM":
        lower = function.lower()
        category = "PA" if "pa" in lower else "LNA" if "lna" in lower else "SWITCH" if "switch" in lower or "sw" in lower else "FEM"
    specs = {
        key: _csv_number(value) if key not in {"digital_interface", "note"} else value
        for key, value in row.items()
        if key not in COMMON_CSV_FIELDS
        and value not in (None, "")
    }
    return Component(
        component_id=None,
        category=category,
        manufacturer=str(row.get("manufacturer") or "").strip(),
        part_no=str(row.get("part_no") or "").strip(),
        application=row.get("application") or "Common",
        function=function,
        process=row.get("process") or None,
        freq_min_ghz=_csv_number(row.get("freq_min_ghz")),
        freq_max_ghz=_csv_number(row.get("freq_max_ghz")),
        supply_voltage_v=_csv_number(row.get("supply_voltage_v")),
        power_consumption_w=_csv_number(row.get("power_consumption_w")),
        package=row.get("package") or None,
        operating_temp_min_c=_csv_number(row.get("operating_temp_min_c")),
        operating_temp_max_c=_csv_number(row.get("operating_temp_max_c")),
        product_url=row.get("product_url") or None,
        datasheet_url=row.get("datasheet_url") or None,
        datasheet_revision=row.get("datasheet_revision") or None,
        datasheet_page=row.get("datasheet_page") or None,
        source_file=row.get("source_file") or source_file,
        source_date=row.get("source_date") or row.get("retrieved_at") or None,
        extraction_method=row.get("extraction_method") or None,
        extraction_evidence=row.get("extraction_evidence") or None,
        note=row.get("note") or None,
        image_url=row.get("image_url") or None,
        specs=specs,
    )


def import_components_csv(path: Path, repository: ComponentRepository) -> int:
    count = 0
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if not row.get("part_no") or not row.get("manufacturer"): continue
            component = component_from_csv_row(row, Path(path).name)
            repository.upsert(component)
            count += 1
    return count
